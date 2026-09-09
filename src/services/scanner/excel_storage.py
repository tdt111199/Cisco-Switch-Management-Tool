"""Persistent Excel Storage Manager for Cisco Historical Unused Port Scanner.

Manages data/port_history.xlsx with 4 persistent sheets:
- Port_History: Complete snapshot of every scan run (append-only)
- Port_Summary: Aggregated current state & historical intelligence per port
- Protected_Ports: List of ports permanently protected from shutdown
- Scan_Log: History of execution scans and metrics

Features automatic initialization, automatic pre-write backups, and safe atomic saving.
"""

import os
import shutil
import time
from datetime import datetime
from typing import Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.logger import logger
from core.scanner_models import (
    PortClassification,
    PortSnapshot,
    PortSummaryRecord,
    ProtectedPortRecord,
    ScanLogRecord,
)


class ExcelStorageError(Exception):
    """Custom exception for Excel storage errors."""
    pass


class ExcelStorageManager:
    """Manages reading, writing, and backing up port_history.xlsx."""

    SHEET_HISTORY = "Port_History"
    SHEET_SUMMARY = "Port_Summary"
    SHEET_PROTECTED = "Protected_Ports"
    SHEET_SCAN_LOG = "Scan_Log"

    HEADER_STYLE = {
        "fill": PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }

    def __init__(self, file_path: str = "data/port_history.xlsx", backup_dir: str = "backups/excel_history"):
        self.file_path = os.path.abspath(file_path)
        self.backup_dir = os.path.abspath(backup_dir)
        self.ensure_initialized()

    def ensure_initialized(self) -> None:
        """Ensure file and all 4 sheets exist with proper headers."""
        os.makedirs(os.path.dirname(self.file_path), exist_ok=True)
        os.makedirs(self.backup_dir, exist_ok=True)

        if not os.path.exists(self.file_path):
            logger.info(f"Tệp Excel lưu trữ chưa tồn tại. Đang tự động khởi tạo: {self.file_path}")
            wb = openpyxl.Workbook()
            # Remove default sheet
            wb.remove(wb.active)

            self._create_history_sheet(wb)
            self._create_summary_sheet(wb)
            self._create_protected_sheet(wb)
            self._create_scan_log_sheet(wb)

            self._atomic_save(wb)
            logger.success("Đã khởi tạo thành công tệp Excel lưu trữ port_history.xlsx.")
        else:
            # Check if all sheets exist; if any missing, create them
            try:
                wb = openpyxl.load_workbook(self.file_path)
                modified = False
                if self.SHEET_HISTORY not in wb.sheetnames:
                    self._create_history_sheet(wb)
                    modified = True
                if self.SHEET_SUMMARY not in wb.sheetnames:
                    self._create_summary_sheet(wb)
                    modified = True
                if self.SHEET_PROTECTED not in wb.sheetnames:
                    self._create_protected_sheet(wb)
                    modified = True
                if self.SHEET_SCAN_LOG not in wb.sheetnames:
                    self._create_scan_log_sheet(wb)
                    modified = True

                if modified:
                    self._atomic_save(wb)
            except Exception as e:
                logger.error(f"Lỗi kiểm tra tệp Excel lưu trữ: {e}")

    def backup_storage_file(self) -> Optional[str]:
        """Creates a timestamped backup of the current Excel file before modification."""
        if not os.path.exists(self.file_path):
            return None
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_name = f"port_history_backup_{timestamp}.xlsx"
            backup_path = os.path.join(self.backup_dir, backup_name)
            shutil.copy2(self.file_path, backup_path)
            logger.info(f"Đã sao lưu tự động tệp lịch sử ra: {backup_path}")
            return backup_path
        except Exception as e:
            logger.error(f"Lỗi khi sao lưu tệp Excel lịch sử: {e}")
            return None

    def _atomic_save(self, wb: openpyxl.Workbook) -> None:
        """Saves workbook atomically using a temp file to prevent file corruption."""
        temp_file = f"{self.file_path}.tmp_{int(time.time()*1000)}"
        try:
            wb.save(temp_file)
            wb.close()
            # Replace target with temp file
            if os.path.exists(self.file_path):
                os.replace(temp_file, self.file_path)
            else:
                shutil.move(temp_file, self.file_path)
        except PermissionError:
            if os.path.exists(temp_file):
                try:
                    os.remove(temp_file)
                except Exception:
                    pass
            msg = (
                f"Không thể ghi tệp Excel vì tệp đang được mở bởi ứng dụng khác (Microsoft Excel).\n"
                f"Vui lòng đóng tệp '{os.path.basename(self.file_path)}' và thử lại!"
            )
            logger.error(msg)
            raise ExcelStorageError(msg)
        except Exception as e:
            if os.path.exists(temp_file):
                try:
                    os.remove(temp_file)
                except Exception:
                    pass
            msg = f"Lỗi không xác định khi lưu tệp Excel: {e}"
            logger.error(msg)
            raise ExcelStorageError(msg)

    # ------------------ Sheet Creation Helpers ------------------ #

    def _create_history_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title=self.SHEET_HISTORY)
        headers = [
            "Scan_Time", "Switch_Name", "Switch_IP", "Interface", "Status",
            "Admin_Status", "VLAN", "Description", "MAC_Addresses", "In_Packets",
            "Out_Packets", "Last_Input", "Last_Output", "Last_Activity_Days",
            "Is_Trunk", "Switchport_Mode", "Is_PortChannel", "PortChannel_ID",
            "CDP_Neighbor", "LLDP_Neighbor", "Neighbor_Info", "Speed", "Duplex", "Raw_Info"
        ]
        self._apply_headers(ws, headers)

    def _create_summary_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title=self.SHEET_SUMMARY)
        headers = [
            "Switch_Name", "Switch_IP", "Interface", "Current_Status", "Current_VLAN",
            "Description", "Current_MAC", "Speed_Duplex", "Classification",
            "Days_Unused", "Consecutive_Unused_Scans", "First_Seen", "Last_Seen",
            "Last_Active_Time", "Active_Transitions", "Is_Protected",
            "Protection_Reason", "Classification_Reason", "Last_Input", "Last_Output"
        ]
        self._apply_headers(ws, headers)

    def _create_protected_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title=self.SHEET_PROTECTED)
        headers = ["Switch_Name", "Interface", "Reason", "Added_Date", "Added_By"]
        self._apply_headers(ws, headers)

    def _create_scan_log_sheet(self, wb: openpyxl.Workbook):
        ws = wb.create_sheet(title=self.SHEET_SCAN_LOG)
        headers = [
            "Scan_ID", "Timestamp", "Switches_Scanned", "Total_Ports", "Active_Count",
            "Unused_Count", "Monitor_Count", "Protected_Count", "Trunk_Count",
            "Error_Count", "Duration_Seconds", "Status", "Notes"
        ]
        self._apply_headers(ws, headers)

    def _apply_headers(self, ws, headers: List[str]):
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = self.HEADER_STYLE["fill"]
            cell.font = self.HEADER_STYLE["font"]
            cell.alignment = self.HEADER_STYLE["alignment"]
            # Set estimated width
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = max(len(header) + 4, 14)

    # ------------------ Reading Methods ------------------ #

    def load_protected_ports(self) -> List[ProtectedPortRecord]:
        """Loads all protected ports from Protected_Ports sheet."""
        records = []
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if self.SHEET_PROTECTED not in wb.sheetnames:
                return records
            ws = wb[self.SHEET_PROTECTED]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0] or not row[1]:
                    continue
                records.append(ProtectedPortRecord(
                    switch_name=str(row[0]).strip(),
                    interface=str(row[1]).strip(),
                    reason=str(row[2]).strip() if len(row) > 2 and row[2] else "User Protected",
                    added_date=str(row[3]).strip() if len(row) > 3 and row[3] else "",
                    added_by=str(row[4]).strip() if len(row) > 4 and row[4] else "Admin"
                ))
            wb.close()
        except Exception as e:
            logger.error(f"Lỗi khi đọc Protected_Ports từ Excel: {e}")
        return records

    def load_port_summaries(self) -> List[PortSummaryRecord]:
        """Loads all aggregated records from Port_Summary sheet."""
        records = []
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if self.SHEET_SUMMARY not in wb.sheetnames:
                return records
            ws = wb[self.SHEET_SUMMARY]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0] or not row[2]:
                    continue
                # Map row fields
                sw_name = str(row[0] or "").strip()
                sw_ip = str(row[1] or "").strip()
                iface = str(row[2] or "").strip()
                status = str(row[3] or "").strip()
                vlan = str(row[4] or "").strip()
                desc = str(row[5] or "").strip()
                macs_str = str(row[6] or "").strip()
                macs = [m.strip() for m in macs_str.split(",") if m.strip()] if macs_str else []
                spd_dup = str(row[7] or "").strip()
                cls_str = str(row[8] or "").strip().upper()
                try:
                    cls_val = PortClassification(cls_str)
                except ValueError:
                    cls_val = PortClassification.MONITOR

                days_unused = int(row[9]) if len(row) > 9 and row[9] is not None and str(row[9]).isdigit() else 0
                consec_scans = int(row[10]) if len(row) > 10 and row[10] is not None and str(row[10]).isdigit() else 0
                first_seen = str(row[11] or "").strip() if len(row) > 11 else ""
                last_seen = str(row[12] or "").strip() if len(row) > 12 else ""
                last_active = str(row[13] or "").strip() if len(row) > 13 else ""
                transitions = int(row[14]) if len(row) > 14 and row[14] is not None and str(row[14]).isdigit() else 0
                is_prot = str(row[15]).strip().lower() in ("true", "1", "yes") if len(row) > 15 and row[15] is not None else False
                prot_reason = str(row[16] or "").strip() if len(row) > 16 else ""
                cls_reason = str(row[17] or "").strip() if len(row) > 17 else ""
                last_in = str(row[18] or "").strip() if len(row) > 18 else ""
                last_out = str(row[19] or "").strip() if len(row) > 19 else ""

                records.append(PortSummaryRecord(
                    switch_name=sw_name,
                    switch_ip=sw_ip,
                    interface=iface,
                    current_status=status,
                    current_vlan=vlan,
                    description=desc,
                    current_macs=macs,
                    speed_duplex=spd_dup,
                    classification=cls_val,
                    days_unused=days_unused,
                    consecutive_unused_scans=consec_scans,
                    first_seen=first_seen,
                    last_seen=last_seen,
                    last_active_time=last_active,
                    active_transitions=transitions,
                    is_protected=is_prot,
                    protection_reason=prot_reason,
                    classification_reason=cls_reason,
                    last_input_str=last_in,
                    last_output_str=last_out,
                ))
            wb.close()
        except Exception as e:
            logger.error(f"Lỗi khi đọc Port_Summary từ Excel: {e}")
        return records

    def load_scan_history(self, limit: Optional[int] = None) -> List[PortSnapshot]:
        """Loads historical snapshots from Port_History."""
        snapshots = []
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if self.SHEET_HISTORY not in wb.sheetnames:
                return snapshots
            ws = wb[self.SHEET_HISTORY]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            if limit:
                rows = rows[-limit:]

            for row in rows:
                if not row or not row[0] or not row[1]:
                    continue
                macs_str = str(row[8] or "")
                macs = [m.strip() for m in macs_str.split(",") if m.strip()] if macs_str else []
                snapshots.append(PortSnapshot(
                    scan_time=str(row[0] or ""),
                    switch_name=str(row[1] or ""),
                    switch_ip=str(row[2] or ""),
                    interface=str(row[3] or ""),
                    status=str(row[4] or ""),
                    admin_status=str(row[5] or ""),
                    vlan=str(row[6] or ""),
                    description=str(row[7] or ""),
                    mac_addresses=macs,
                    in_packets=int(row[9] or 0),
                    out_packets=int(row[10] or 0),
                    last_input_str=str(row[11] or ""),
                    last_output_str=str(row[12] or ""),
                    last_activity_days=float(row[13] or -1.0),
                    is_trunk=str(row[14]).lower() in ("true", "1", "yes"),
                    switchport_mode=str(row[15] or ""),
                    is_portchannel=str(row[16]).lower() in ("true", "1", "yes"),
                    port_channel_id=str(row[17] or ""),
                    has_cdp_neighbor=str(row[18]).lower() in ("true", "1", "yes"),
                    has_lldp_neighbor=str(row[19]).lower() in ("true", "1", "yes"),
                    neighbor_info=str(row[20] or ""),
                    speed=str(row[21] or ""),
                    duplex=str(row[22] or ""),
                    raw_info=str(row[23] or "") if len(row) > 23 else ""
                ))
            wb.close()
        except Exception as e:
            logger.error(f"Lỗi khi đọc Port_History từ Excel: {e}")
        return snapshots

    def load_scan_logs(self) -> List[ScanLogRecord]:
        """Loads execution scan logs."""
        logs = []
        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if self.SHEET_SCAN_LOG not in wb.sheetnames:
                return logs
            ws = wb[self.SHEET_SCAN_LOG]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                logs.append(ScanLogRecord(
                    scan_id=str(row[0] or ""),
                    timestamp=str(row[1] or ""),
                    switches_scanned=int(row[2] or 0),
                    total_ports=int(row[3] or 0),
                    active_count=int(row[4] or 0),
                    unused_count=int(row[5] or 0),
                    monitor_count=int(row[6] or 0),
                    protected_count=int(row[7] or 0),
                    trunk_count=int(row[8] or 0),
                    error_count=int(row[9] or 0),
                    duration_seconds=float(row[10] or 0.0),
                    status=str(row[11] or "SUCCESS"),
                    notes=str(row[12] or "") if len(row) > 12 else ""
                ))
            wb.close()
        except Exception as e:
            logger.error(f"Lỗi khi đọc Scan_Log từ Excel: {e}")
        return logs

    # ------------------ Writing & Persistence Methods ------------------ #

    def save_scan_results(
        self,
        new_snapshots: List[PortSnapshot],
        updated_summaries: List[PortSummaryRecord],
        scan_log: ScanLogRecord,
    ) -> None:
        """Appends new snapshots, updates summary sheet, and appends scan log.

        Performs automatic backup beforehand and uses atomic save.
        """
        # 1. Backup existing storage file
        self.backup_storage_file()

        # 2. Open workbook
        try:
            wb = openpyxl.load_workbook(self.file_path)
        except Exception as e:
            logger.error(f"Không thể đọc file {self.file_path} để cập nhật: {e}. Đang tái tạo...")
            self.ensure_initialized()
            wb = openpyxl.load_workbook(self.file_path)

        # 3. Append to Port_History
        if self.SHEET_HISTORY not in wb.sheetnames:
            self._create_history_sheet(wb)
        ws_hist = wb[self.SHEET_HISTORY]
        for s in new_snapshots:
            ws_hist.append([
                s.scan_time,
                s.switch_name,
                s.switch_ip,
                s.interface,
                s.status,
                s.admin_status,
                s.vlan,
                s.description,
                ", ".join(s.mac_addresses),
                s.in_packets,
                s.out_packets,
                s.last_input_str,
                s.last_output_str,
                s.last_activity_days,
                "Yes" if s.is_trunk else "No",
                s.switchport_mode,
                "Yes" if s.is_portchannel else "No",
                s.port_channel_id,
                "Yes" if s.has_cdp_neighbor else "No",
                "Yes" if s.has_lldp_neighbor else "No",
                s.neighbor_info,
                s.speed,
                s.duplex,
                s.raw_info
            ])

        # 4. Overwrite Port_Summary with updated records
        if self.SHEET_SUMMARY in wb.sheetnames:
            wb.remove(wb[self.SHEET_SUMMARY])
        self._create_summary_sheet(wb)
        ws_sum = wb[self.SHEET_SUMMARY]

        # Classification row styling
        fill_unused = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # soft red
        fill_active = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # soft green
        fill_monitor = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # soft yellow
        fill_prot = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")    # soft sky blue
        fill_trunk = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid")   # soft purple

        for r in updated_summaries:
            row_data = [
                r.switch_name,
                r.switch_ip,
                r.interface,
                r.current_status,
                r.current_vlan,
                r.description,
                ", ".join(r.current_macs),
                r.speed_duplex,
                r.classification.value if hasattr(r.classification, "value") else str(r.classification),
                r.days_unused,
                r.consecutive_unused_scans,
                r.first_seen,
                r.last_seen,
                r.last_active_time,
                r.active_transitions,
                "Yes" if r.is_protected else "No",
                r.protection_reason,
                r.classification_reason,
                r.last_input_str,
                r.last_output_str
            ]
            ws_sum.append(row_data)
            cur_row = ws_sum.max_row
            cls_cell = ws_sum.cell(row=cur_row, column=9)
            if r.classification == PortClassification.UNUSED:
                cls_cell.fill = fill_unused
                cls_cell.font = Font(color="B91C1C", bold=True)
            elif r.classification == PortClassification.ACTIVE:
                cls_cell.fill = fill_active
                cls_cell.font = Font(color="15803D", bold=True)
            elif r.classification == PortClassification.MONITOR:
                cls_cell.fill = fill_monitor
                cls_cell.font = Font(color="B45309", bold=True)
            elif r.classification == PortClassification.PROTECTED:
                cls_cell.fill = fill_prot
                cls_cell.font = Font(color="0369A1", bold=True)
            elif r.classification in (PortClassification.TRUNK_UPLINK, PortClassification.PORT_CHANNEL):
                cls_cell.fill = fill_trunk
                cls_cell.font = Font(color="6B21A8", bold=True)

        # 5. Append to Scan_Log
        if self.SHEET_SCAN_LOG not in wb.sheetnames:
            self._create_scan_log_sheet(wb)
        ws_log = wb[self.SHEET_SCAN_LOG]
        ws_log.append([
            scan_log.scan_id,
            scan_log.timestamp,
            scan_log.switches_scanned,
            scan_log.total_ports,
            scan_log.active_count,
            scan_log.unused_count,
            scan_log.monitor_count,
            scan_log.protected_count,
            scan_log.trunk_count,
            scan_log.error_count,
            round(scan_log.duration_seconds, 2),
            scan_log.status,
            scan_log.notes
        ])

        # 6. Save atomically
        self._atomic_save(wb)
        logger.success(f"Đã lưu thành công dữ liệu phiên quét vào {self.file_path}")

    def update_protected_ports(self, protected_records: List[ProtectedPortRecord]) -> None:
        """Overwrites Protected_Ports sheet with current protected list."""
        self.backup_storage_file()
        try:
            wb = openpyxl.load_workbook(self.file_path)
            if self.SHEET_PROTECTED in wb.sheetnames:
                wb.remove(wb[self.SHEET_PROTECTED])
            self._create_protected_sheet(wb)
            ws = wb[self.SHEET_PROTECTED]

            for r in protected_records:
                ws.append([r.switch_name, r.interface, r.reason, r.added_date, r.added_by])

            self._atomic_save(wb)
            logger.info("Đã cập nhật danh sách Protected_Ports trong Excel.")
        except Exception as e:
            logger.error(f"Lỗi khi lưu Protected_Ports vào Excel: {e}")
            raise e
