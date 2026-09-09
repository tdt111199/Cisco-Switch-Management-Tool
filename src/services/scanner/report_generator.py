"""Report Generator for Cisco Historical Unused Port Scanner.

Exports formatted, executive-ready Excel (.xlsx) and CSV reports with KPI metrics,
switch breakdown table with unused percentages, detailed port lists, and unused-only sheets.
"""

import csv
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.logger import logger
from core.scanner_models import PortClassification, PortSummaryRecord


class ReportGenerator:
    """Generates Excel and CSV reports."""

    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    THIN_BORDER = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    @classmethod
    def export_excel_report(cls, summaries: List[PortSummaryRecord], file_path: str) -> None:
        """Exports a full multi-sheet audit report in Excel format."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # 1. Sheet: Dashboard Tổng Hợp
        cls._build_dashboard_sheet(wb, summaries)

        # 2. Sheet: Toàn Bộ Cổng
        cls._build_ports_sheet(wb, "Tất Cả Cổng", summaries)

        # 3. Sheet: Cổng UNUSED
        unused_ports = [p for p in summaries if p.classification == PortClassification.UNUSED]
        cls._build_ports_sheet(wb, "Cổng UNUSED Đủ Điều Kiện", unused_ports)

        wb.save(file_path)
        wb.close()
        logger.success(f"Đã xuất báo cáo Excel thành công: {file_path}")

    @classmethod
    def _build_dashboard_sheet(cls, wb: openpyxl.Workbook, summaries: List[PortSummaryRecord]):
        ws = wb.create_sheet(title="Dashboard Tổng Hợp")
        ws.views.sheetView[0].showGridLines = True

        # Title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "BÁO CÁO TỔNG HỢP LỊCH SỬ SỬ DỤNG PORT SWITCH CISCO"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        ws.merge_cells("A2:G2")
        ws["A2"].value = f"Ngày xuất báo cáo: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="4B5563")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Global KPI calculations
        switches_set = {p.switch_name for p in summaries}
        total_ports = len(summaries)
        active_count = sum(1 for p in summaries if p.classification == PortClassification.ACTIVE)
        unused_count = sum(1 for p in summaries if p.classification == PortClassification.UNUSED)
        monitor_count = sum(1 for p in summaries if p.classification == PortClassification.MONITOR)
        protected_count = sum(1 for p in summaries if p.classification == PortClassification.PROTECTED)
        trunk_count = sum(1 for p in summaries if p.classification in (PortClassification.TRUNK_UPLINK, PortClassification.PORT_CHANNEL))

        # KPI Block
        kpi_labels = ["Tổng Switch", "Tổng Số Port", "Port ACTIVE", "Port UNUSED", "Port MONITOR", "Port PROTECTED", "TRUNK / PO"]
        kpi_vals = [len(switches_set), total_ports, active_count, unused_count, monitor_count, protected_count, trunk_count]
        kpi_colors = ["E0F2FE", "F3F4F6", "DCFCE7", "FEE2E2", "FEF9C3", "DBEAFE", "F3E8FF"]
        kpi_text_colors = ["0369A1", "374151", "15803D", "B91C1C", "B45309", "1D4ED8", "6B21A8"]

        ws.row_dimensions[4].height = 20
        ws.row_dimensions[5].height = 28

        for col_idx, (lbl, val, bg_col, txt_col) in enumerate(zip(kpi_labels, kpi_vals, kpi_colors, kpi_text_colors), start=1):
            c_lbl = ws.cell(row=4, column=col_idx, value=lbl)
            c_lbl.font = Font(name="Calibri", size=10, bold=True, color="4B5563")
            c_lbl.alignment = Alignment(horizontal="center", vertical="center")
            c_lbl.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
            c_lbl.border = cls.THIN_BORDER

            c_val = ws.cell(row=5, column=col_idx, value=val)
            c_val.font = Font(name="Calibri", size=16, bold=True, color=txt_col)
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
            c_val.border = cls.THIN_BORDER
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 18

        # Section Header: Switch Breakdown
        ws.cell(row=7, column=1, value="CHI TIẾT THEO TỪNG THIẾT BỊ SWITCH").font = Font(size=12, bold=True, color="1E3A8A")
        ws.row_dimensions[7].height = 24

        table_headers = [
            "Tên Switch", "Địa Chỉ IP", "Tổng Số Port", "Port ACTIVE",
            "Port UNUSED", "Port MONITOR", "Port PROTECTED", "Tỷ Lệ UNUSED (%)"
        ]
        ws.row_dimensions[8].height = 25
        for col_idx, th in enumerate(table_headers, start=1):
            cell = ws.cell(row=8, column=col_idx, value=th)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        # Group data by switch
        by_switch: Dict[Tuple[str, str], List[PortSummaryRecord]] = {}
        for p in summaries:
            key = (p.switch_name, p.switch_ip)
            if key not in by_switch:
                by_switch[key] = []
            by_switch[key].append(p)

        row_curr = 9
        for (sw_name, sw_ip), sw_ports in sorted(by_switch.items(), key=lambda x: x[0][0]):
            tot = len(sw_ports)
            act = sum(1 for p in sw_ports if p.classification == PortClassification.ACTIVE)
            un = sum(1 for p in sw_ports if p.classification == PortClassification.UNUSED)
            mon = sum(1 for p in sw_ports if p.classification == PortClassification.MONITOR)
            prot = sum(1 for p in sw_ports if p.classification == PortClassification.PROTECTED)
            rate = round((un / tot * 100.0), 1) if tot > 0 else 0.0

            ws.row_dimensions[row_curr].height = 20
            row_vals = [sw_name, sw_ip, tot, act, un, mon, prot, f"{rate}%"]
            for col_idx, v in enumerate(row_vals, start=1):
                cell = ws.cell(row=row_curr, column=col_idx, value=v)
                cell.font = Font(name="Calibri", size=10)
                cell.border = cls.THIN_BORDER
                if col_idx in (1, 2):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Highlight unused rate
                if col_idx == 8:
                    cell.font = Font(name="Calibri", size=10, bold=True)
                    if rate >= 30.0:
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        cell.font = Font(color="B91C1C", bold=True)

            row_curr += 1

    @classmethod
    def _build_ports_sheet(cls, wb: openpyxl.Workbook, sheet_title: str, ports: List[PortSummaryRecord]):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Tên Switch", "IP Switch", "Interface", "Trạng Thái", "VLAN",
            "Mô Tả (Description)", "Địa Chỉ MAC", "Speed/Duplex", "Phân Loại",
            "Số Ngày Inactive", "Chuỗi Scan Inactive", "First Seen", "Last Seen",
            "Last Active Time", "Bảo Vệ?", "Lý Do Phân Loại"
        ]
        ws.row_dimensions[1].height = 28
        for col_idx, th in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=th)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(len(th) + 4, 15)

        # Fills
        fill_un = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        fill_act = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fill_mon = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        fill_prot = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
        fill_trunk = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid")

        for r_idx, p in enumerate(ports, start=2):
            ws.row_dimensions[r_idx].height = 20
            row_data = [
                p.switch_name,
                p.switch_ip,
                p.interface,
                p.current_status,
                p.current_vlan,
                p.description,
                ", ".join(p.current_macs),
                p.speed_duplex,
                p.classification.value if hasattr(p.classification, "value") else str(p.classification),
                p.days_unused,
                p.consecutive_unused_scans,
                p.first_seen,
                p.last_seen,
                p.last_active_time,
                "Có" if p.is_protected else "Không",
                p.classification_reason
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name="Calibri", size=10)
                cell.border = cls.THIN_BORDER
                if c_idx in (1, 2, 3, 6, 16):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Classification badge
                if c_idx == 9:
                    if p.classification == PortClassification.UNUSED:
                        cell.fill = fill_un
                        cell.font = Font(color="B91C1C", bold=True)
                    elif p.classification == PortClassification.ACTIVE:
                        cell.fill = fill_act
                        cell.font = Font(color="15803D", bold=True)
                    elif p.classification == PortClassification.MONITOR:
                        cell.fill = fill_mon
                        cell.font = Font(color="B45309", bold=True)
                    elif p.classification == PortClassification.PROTECTED:
                        cell.fill = fill_prot
                        cell.font = Font(color="0369A1", bold=True)
                    elif p.classification in (PortClassification.TRUNK_UPLINK, PortClassification.PORT_CHANNEL):
                        cell.fill = fill_trunk
                        cell.font = Font(color="6B21A8", bold=True)

    @classmethod
    def export_csv_report(cls, summaries: List[PortSummaryRecord], file_path: str) -> None:
        """Exports CSV formatted report with UTF-8 BOM."""
        fieldnames = [
            "Switch_Name", "Switch_IP", "Interface", "Status", "VLAN",
            "Description", "MAC_Addresses", "Speed_Duplex", "Classification",
            "Days_Unused", "Consecutive_Unused_Scans", "First_Seen", "Last_Seen",
            "Last_Active_Time", "Is_Protected", "Classification_Reason"
        ]

        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for p in summaries:
                writer.writerow({
                    "Switch_Name": p.switch_name,
                    "Switch_IP": p.switch_ip,
                    "Interface": p.interface,
                    "Status": p.current_status,
                    "VLAN": p.current_vlan,
                    "Description": p.description,
                    "MAC_Addresses": ", ".join(p.current_macs),
                    "Speed_Duplex": p.speed_duplex,
                    "Classification": p.classification.value if hasattr(p.classification, "value") else str(p.classification),
                    "Days_Unused": p.days_unused,
                    "Consecutive_Unused_Scans": p.consecutive_unused_scans,
                    "First_Seen": p.first_seen,
                    "Last_Seen": p.last_seen,
                    "Last_Active_Time": p.last_active_time,
                    "Is_Protected": "Yes" if p.is_protected else "No",
                    "Classification_Reason": p.classification_reason
                })

        logger.success(f"Đã xuất báo cáo CSV thành công: {file_path}")
