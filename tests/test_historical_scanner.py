"""Comprehensive unit tests for Cisco Historical Unused Port Scanner modules."""

import os
import shutil
import sys
import unittest
from datetime import datetime

# Ensure src/ and root are in sys.path
sys.path.insert(0, os.path.abspath("src"))
sys.path.insert(0, os.path.abspath("."))

from core.inventory import SwitchDevice
from core.scanner_models import (
    PortClassification,
    PortSnapshot,
    PortSummaryRecord,
    ProtectedPortRecord,
    ScanLogRecord,
    ScanSettings,
)
from services.scanner.cisco_collector import (
    CiscoDataCollector,
    canonical_short_name,
    normalize_interface_name,
    parse_cisco_duration,
)
from services.scanner.config_generator import ConfigGenerator
from services.scanner.excel_storage import ExcelStorageManager
from services.scanner.historical_analyzer import HistoricalAnalyzer
from services.scanner.report_generator import ReportGenerator
from services.scanner.risk_protection_engine import RiskProtectionEngine


class TestCiscoCollectorParsing(unittest.TestCase):
    """Tests duration parser and CLI text parser logic."""

    def test_duration_parser(self):
        self.assertEqual(parse_cisco_duration("never"), 9999.0)
        self.assertEqual(parse_cisco_duration("12w3d"), 12 * 7 + 3)  # 87 days
        self.assertAlmostEqual(parse_cisco_duration("3d12h"), 3.5, places=1)
        self.assertAlmostEqual(parse_cisco_duration("00:30:00"), 0.02, places=2)
        self.assertGreater(parse_cisco_duration("1y"), 360)

    def test_interface_normalization(self):
        self.assertEqual(normalize_interface_name("Gi0/1"), "GigabitEthernet0/1")
        self.assertEqual(normalize_interface_name("Fa0/24"), "FastEthernet0/24")
        self.assertEqual(normalize_interface_name("Te1/0/1"), "TenGigabitEthernet1/0/1")
        self.assertEqual(normalize_interface_name("Po1"), "Port-channel1")
        self.assertEqual(canonical_short_name("GigabitEthernet0/1"), "Gi0/1")
        self.assertEqual(canonical_short_name("Port-channel2"), "Po2")

    def test_mock_cli_parsing(self):
        dev = SwitchDevice(name="SW-Test-01", ip="192.168.1.10", username="admin")
        collector = CiscoDataCollector(dev)

        raw_status = (
            "Port      Name               Status       Vlan       Duplex  Speed Type\n"
            "Gi0/1     Uplink to Core     connected    trunk        a-full  a-1000 10/100/1000BaseTX\n"
            "Gi0/2                        notconnect   10           auto    auto 10/100/1000BaseTX\n"
            "Gi0/3     Printer HR         connected    20         a-full  a-100  10/100/1000BaseTX\n"
            "Po1                          connected    trunk      a-full  a-1000\n"
        )
        raw_desc = (
            "Interface                      Status         Protocol Description\n"
            "Gi0/1                          up             up       Uplink to Core\n"
            "Gi0/2                          down           down     \n"
            "Gi0/3                          up             up       Printer HR\n"
        )
        raw_switchport = (
            "Name: Gi0/1\n"
            "Operational Mode: trunk\n"
            "Name: Gi0/2\n"
            "Operational Mode: static access\n"
            "Access Mode VLAN: 10\n"
        )
        raw_etherchannel = (
            "Group  Port-channel  Protocol    Ports\n"
            "------+-------------+-----------+-----------------------------------------------\n"
            "1      Po1(SU)         LACP      Gi0/4(P)    Gi0/5(P)\n"
        )
        raw_mac = (
            "Vlan    Mac Address       Type        Ports\n"
            "----    -----------       --------    -----\n"
            "  10    0011.2233.4455    DYNAMIC     Gi0/3\n"
        )
        raw_cdp = (
            "Device ID        Local Intrfce     Holdtme    Capability  Platform  Port ID\n"
            "Core-SW          Gig 0/1           150              R S I WS-C3850  Gig 1/0/1\n"
        )
        raw_interfaces = (
            "GigabitEthernet0/1 is up, line protocol is up\n"
            "  Last input 00:00:01, output 00:00:00\n"
            "  100000 packets input, 50000000 bytes\n"
            "GigabitEthernet0/2 is down, line protocol is down\n"
            "  Last input 14w2d, output 14w2d\n"
            "  0 packets input, 0 bytes\n"
        )

        snaps = collector._parse_all_data(
            scan_time="2026-09-09 16:00:00",
            raw_status=raw_status,
            raw_desc=raw_desc,
            raw_switchport=raw_switchport,
            raw_etherchannel=raw_etherchannel,
            raw_mac=raw_mac,
            raw_cdp=raw_cdp,
            raw_lldp="",
            raw_interfaces=raw_interfaces,
        )

        snap_map = {s.interface: s for s in snaps}
        self.assertIn("GigabitEthernet0/1", snap_map)
        self.assertIn("GigabitEthernet0/2", snap_map)
        self.assertIn("GigabitEthernet0/3", snap_map)

        # Gi0/1 should be trunk & have CDP neighbor
        g1 = snap_map["GigabitEthernet0/1"]
        self.assertTrue(g1.is_trunk)
        self.assertTrue(g1.has_cdp_neighbor)

        # Gi0/2 should be notconnect and have last activity around 100 days
        g2 = snap_map["GigabitEthernet0/2"]
        self.assertEqual(g2.status, "notconnect")
        self.assertGreater(g2.last_activity_days, 90)

        # Gi0/3 should have MAC learned
        g3 = snap_map["GigabitEthernet0/3"]
        self.assertEqual(len(g3.mac_addresses), 1)


class TestExcelStorage(unittest.TestCase):
    """Tests Excel persistent storage auto-initialization, backup, and sheets."""

    def setUp(self):
        self.test_dir = "tests/test_storage"
        self.test_excel = os.path.join(self.test_dir, "port_history_test.xlsx")
        self.backup_dir = os.path.join(self.test_dir, "backups")
        os.makedirs(self.test_dir, exist_ok=True)
        self.storage = ExcelStorageManager(file_path=self.test_excel, backup_dir=self.backup_dir)

    def tearDown(self):
        if os.path.exists(self.test_dir):
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_file_and_sheets_initialized(self):
        self.assertTrue(os.path.exists(self.test_excel))
        import openpyxl
        wb = openpyxl.load_workbook(self.test_excel)
        self.assertIn("Port_History", wb.sheetnames)
        self.assertIn("Port_Summary", wb.sheetnames)
        self.assertIn("Protected_Ports", wb.sheetnames)
        self.assertIn("Scan_Log", wb.sheetnames)
        wb.close()

    def test_save_and_backup(self):
        snap = PortSnapshot(
            scan_time="2026-09-09 10:00:00",
            switch_name="SW-01",
            switch_ip="10.0.0.1",
            interface="GigabitEthernet0/1",
            status="notconnect",
            admin_status="up",
            vlan="10",
        )
        summary = PortSummaryRecord(
            switch_name="SW-01",
            switch_ip="10.0.0.1",
            interface="GigabitEthernet0/1",
            current_status="notconnect",
            current_vlan="10",
            description="",
            classification=PortClassification.UNUSED,
            days_unused=75,
            consecutive_unused_scans=3,
        )
        log_rec = ScanLogRecord(
            scan_id="SCAN_01",
            timestamp="2026-09-09 10:00:00",
            switches_scanned=1,
            total_ports=1,
            active_count=0,
            unused_count=1,
            monitor_count=0,
            protected_count=0,
            trunk_count=0,
            error_count=0,
            duration_seconds=1.5,
            status="SUCCESS",
        )

        self.storage.save_scan_results([snap], [summary], log_rec)

        # Verify backup was created
        backups = os.listdir(self.backup_dir)
        self.assertTrue(len(backups) >= 1)

        # Verify data loaded back
        loaded_summaries = self.storage.load_port_summaries()
        self.assertEqual(len(loaded_summaries), 1)
        self.assertEqual(loaded_summaries[0].interface, "GigabitEthernet0/1")
        self.assertEqual(loaded_summaries[0].classification, PortClassification.UNUSED)

        # Test Protected Ports update
        prot = ProtectedPortRecord(switch_name="SW-01", interface="GigabitEthernet0/1")
        self.storage.update_protected_ports([prot])
        loaded_prot = self.storage.load_protected_ports()
        self.assertEqual(len(loaded_prot), 1)
        self.assertEqual(loaded_prot[0].interface, "GigabitEthernet0/1")


class TestHistoricalAnalyzerAndResetRule(unittest.TestCase):
    """Tests multi-scan metrics and the consecutive inactive reset rule."""

    def test_inactive_and_reset_rule(self):
        analyzer = HistoricalAnalyzer(ScanSettings(threshold_days=60))
        prot_set = set()

        # Scan 1: Port is inactive
        s1 = [PortSnapshot(
            scan_time="2026-01-01 10:00:00",
            switch_name="SW-01",
            switch_ip="10.0.0.1",
            interface="Gi0/1",
            status="notconnect",
            admin_status="up",
            vlan="10",
            last_activity_days=10.0,
        )]
        res1 = analyzer.analyze(s1, [], prot_set)
        self.assertEqual(res1[0].consecutive_unused_scans, 1)

        # Scan 2: Port remains inactive
        s2 = [PortSnapshot(
            scan_time="2026-01-02 10:00:00",
            switch_name="SW-01",
            switch_ip="10.0.0.1",
            interface="Gi0/1",
            status="notconnect",
            admin_status="up",
            vlan="10",
            last_activity_days=11.0,
        )]
        res2 = analyzer.analyze(s2, res1, prot_set)
        self.assertEqual(res2[0].consecutive_unused_scans, 2)

        # Scan 3: Port becomes ACTIVE! -> MUST RESET consecutive scans to 0
        s3 = [PortSnapshot(
            scan_time="2026-01-03 10:00:00",
            switch_name="SW-01",
            switch_ip="10.0.0.1",
            interface="Gi0/1",
            status="connected",
            admin_status="up",
            vlan="10",
            mac_addresses=["0011.2233.4455"],
        )]
        res3 = analyzer.analyze(s3, res2, prot_set)
        self.assertEqual(res3[0].consecutive_unused_scans, 0, "Consecutive unused scans must be reset to 0 upon activation!")
        self.assertEqual(res3[0].days_unused, 0)
        self.assertEqual(res3[0].active_transitions, 1)
        self.assertEqual(res3[0].last_active_time, "2026-01-03 10:00:00")


class TestRiskProtectionAndConfigGeneration(unittest.TestCase):
    """Tests safety classification rules and config generation."""

    def test_classification_logic(self):
        engine = RiskProtectionEngine(ScanSettings(threshold_days=60))
        prot_set = {("SW-01", "Gi0/10")}

        ports = [
            # 0: Inactive 75 days -> UNUSED
            PortSummaryRecord(switch_name="SW-01", switch_ip="10.0.0.1", interface="Gi0/1", current_status="notconnect", current_vlan="10", description="", days_unused=75),
            # 1: Inactive 15 days (< 60) -> MONITOR
            PortSummaryRecord(switch_name="SW-01", switch_ip="10.0.0.1", interface="Gi0/2", current_status="notconnect", current_vlan="10", description="", days_unused=15),
            # 2: Trunk port -> TRUNK/UPLINK
            PortSummaryRecord(switch_name="SW-01", switch_ip="10.0.0.1", interface="Gi0/3", current_status="notconnect", current_vlan="trunk", description="", is_trunk=True, days_unused=100),
            # 3: Port in Protected list -> PROTECTED
            PortSummaryRecord(switch_name="SW-01", switch_ip="10.0.0.1", interface="Gi0/10", current_status="notconnect", current_vlan="10", description="", days_unused=200),
            # 4: Active port -> ACTIVE
            PortSummaryRecord(switch_name="SW-01", switch_ip="10.0.0.1", interface="Gi0/4", current_status="connected", current_vlan="10", description="", current_macs=["aabb.ccdd.eeff"]),
        ]

        evaluated = engine.evaluate_all(ports, prot_set)
        self.assertEqual(evaluated[0].classification, PortClassification.UNUSED)
        self.assertEqual(evaluated[1].classification, PortClassification.MONITOR)
        self.assertEqual(evaluated[2].classification, PortClassification.TRUNK_UPLINK)
        self.assertEqual(evaluated[3].classification, PortClassification.PROTECTED)
        self.assertEqual(evaluated[4].classification, PortClassification.ACTIVE)

        # Config generation test: ONLY port 0 qualifies!
        shutdown_cfg = ConfigGenerator.generate_shutdown_config(evaluated)
        self.assertIn("interface Gi0/1", shutdown_cfg)
        self.assertIn("shutdown", shutdown_cfg)
        self.assertNotIn("interface Gi0/2", shutdown_cfg)
        self.assertNotIn("interface Gi0/3", shutdown_cfg)
        self.assertNotIn("interface Gi0/10", shutdown_cfg)
        self.assertNotIn("interface Gi0/4", shutdown_cfg)

        # Rollback test
        rollback_cfg = ConfigGenerator.generate_rollback_config(evaluated)
        self.assertIn("interface Gi0/1", rollback_cfg)
        self.assertIn("no shutdown", rollback_cfg)


class TestReportGenerator(unittest.TestCase):
    """Tests Excel and CSV report exports."""

    def test_report_export(self):
        out_excel = "tests/test_report.xlsx"
        out_csv = "tests/test_report.csv"
        ports = [
            PortSummaryRecord(
                switch_name="SW-Core",
                switch_ip="10.1.1.1",
                interface="Gi0/1",
                current_status="notconnect",
                current_vlan="10",
                description="Desk Port",
                classification=PortClassification.UNUSED,
                days_unused=90,
            )
        ]
        try:
            ReportGenerator.export_excel_report(ports, out_excel)
            self.assertTrue(os.path.exists(out_excel))

            ReportGenerator.export_csv_report(ports, out_csv)
            self.assertTrue(os.path.exists(out_csv))
        finally:
            if os.path.exists(out_excel):
                os.remove(out_excel)
            if os.path.exists(out_csv):
                os.remove(out_csv)


if __name__ == "__main__":
    unittest.main()
